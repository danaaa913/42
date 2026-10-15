"""Load common competition document formats into page/section-aware units."""

import csv
import json
from html.parser import HTMLParser
from pathlib import Path
from typing import Dict, Iterable, List, Optional

from .cleaning import clean_text, remove_repeated_page_lines
from .models import SourceUnit
from .privacy import filename_privacy_findings, scan_units_privacy
from .security import scan_units, sha256_file


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".html", ".htm", ".json", ".csv", ".docx", ".pptx", ".xlsx"}


class _VisibleHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hidden = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "svg"}:
            self.hidden += 1
        if tag in {"p", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4", "br"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "svg"} and self.hidden:
            self.hidden -= 1
        if tag in {"p", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _relative(path: Path, root: Path) -> str:
    try:
        return path.relative_to(root).as_posix()
    except ValueError:
        return path.name


def _table_markdown(rows) -> str:
    """Render a table as searchable text while preserving columns and headers."""
    values = []
    for row in rows:
        cells = [clean_text("" if cell is None else str(cell)).replace("|", "\\|") for cell in row]
        if any(cells):
            values.append(cells)
    if not values:
        return ""
    width = max(len(row) for row in values)
    values = [row + [""] * (width - len(row)) for row in values]
    header = values[0]
    body = values[1:]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _load_pdf(path: Path, source: str, stats: Optional[Dict] = None) -> List[SourceUnit]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("PDF support requires pdfplumber") from exc
    units = []
    with pdfplumber.open(path) as pdf:
        total_pages = len(pdf.pages)
        for number, page in enumerate(pdf.pages, 1):
            text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            tables = []
            try:
                tables = [_table_markdown(table) for table in page.extract_tables()]
                tables = [table for table in tables if table]
            except Exception:
                tables = []
            if tables:
                text = text + "\n\n" + "\n\n".join(f"[Table]\n{table}" for table in tables)
            if clean_text(text):
                units.append(SourceUnit(text=text, source=source, page=number, document_type="pdf"))
    if total_pages and len(units) < total_pages / 2:
        print(
            f"Warning: {source} yielded text from only {len(units)}/{total_pages} pages; "
            "it may be scanned and require OCR."
        )
    if stats is not None:
        stats.update({"total_pages": total_pages, "pages_with_text": len(units)})
    return units


def _load_docx(path: Path, source: str) -> List[SourceUnit]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX support requires python-docx") from exc
    document = Document(path)
    units, section, paragraphs = [], "", []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            if paragraphs:
                paragraphs.append("")
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading"):
            if paragraphs:
                units.append(SourceUnit("\n".join(paragraphs), source, section=section, document_type="docx"))
                paragraphs = []
            section = text
        paragraphs.append(text)
    if paragraphs:
        units.append(SourceUnit("\n".join(paragraphs), source, section=section, document_type="docx"))
    for index, table in enumerate(document.tables, 1):
        text = _table_markdown([[cell.text for cell in row.cells] for row in table.rows])
        if text:
            units.append(SourceUnit(text, source, section=f"Table {index}", document_type="docx_table"))
    return units


def _load_pptx(path: Path, source: str, stats: Optional[Dict] = None) -> List[SourceUnit]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("PPTX support requires python-pptx") from exc
    presentation = Presentation(path)
    units = []
    for number, slide in enumerate(presentation.slides, 1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                table = _table_markdown([[cell.text for cell in row.cells] for row in shape.table.rows])
                if table:
                    texts.append(f"[Table]\n{table}")
            elif hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
        if texts:
            units.append(SourceUnit("\n".join(texts), source, page=number, document_type="pptx"))
    if stats is not None:
        stats.update({"total_pages": len(presentation.slides), "pages_with_text": len(units)})
    return units


def _load_xlsx(path: Path, source: str) -> List[SourceUnit]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX support requires openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    units = []
    for sheet in workbook.worksheets:
        text = _table_markdown(sheet.iter_rows(values_only=True))
        if text:
            units.append(SourceUnit(text, source, section=sheet.title, document_type="xlsx"))
    return units


def _json_text(value, prefix="") -> List[str]:
    if isinstance(value, dict):
        lines = []
        for key, child in value.items():
            child_prefix = f"{prefix}.{key}" if prefix else str(key)
            lines.extend(_json_text(child, child_prefix))
        return lines
    if isinstance(value, list):
        lines = []
        for index, child in enumerate(value):
            lines.extend(_json_text(child, f"{prefix}[{index}]"))
        return lines
    return [f"{prefix}: {value}"]


def load_file(path: Path, root: Path, stats: Optional[Dict] = None) -> List[SourceUnit]:
    source = _relative(path, root)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path, source, stats)
    if suffix == ".docx":
        return _load_docx(path, source)
    if suffix == ".pptx":
        return _load_pptx(path, source, stats)
    if suffix == ".xlsx":
        return _load_xlsx(path, source)
    if suffix in {".txt", ".md"}:
        return [SourceUnit(path.read_text(encoding="utf-8", errors="replace"), source, document_type=suffix[1:])]
    if suffix in {".html", ".htm"}:
        parser = _VisibleHTML()
        parser.feed(path.read_text(encoding="utf-8", errors="replace"))
        return [SourceUnit("".join(parser.parts), source, document_type="html")]
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        return [SourceUnit("\n".join(_json_text(value)), source, document_type="json")]
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.DictReader(handle)
            rows = [" | ".join(f"{key}: {value}" for key, value in row.items()) for row in reader]
        return [SourceUnit("\n".join(rows), source, document_type="csv")]
    return []


def load_documents(
    directory: Path,
    diagnostics: Optional[Dict] = None,
    quarantine_suspicious: bool = True,
    quarantine_sensitive: bool = True,
) -> List[SourceUnit]:
    if not directory.exists():
        raise FileNotFoundError(f"Documents directory does not exist: {directory}")
    paths = sorted(path for path in directory.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS)
    if not paths:
        raise RuntimeError(f"No supported documents found in {directory}")
    units = []
    failures = []
    files = []
    for path in paths:
        source = _relative(path, directory)
        file_hash = None
        file_size = None
        try:
            file_hash = sha256_file(path)
            file_size = path.stat().st_size
            file_stats = {}
            loaded = load_file(path, directory, file_stats)
            security = scan_units(loaded)
            privacy = scan_units_privacy(loaded)
            filename_findings = filename_privacy_findings(source)
            if filename_findings:
                privacy["findings"] = filename_findings + list(privacy["findings"])
                privacy["high_findings"] = int(privacy["high_findings"]) + sum(
                    item["severity"] == "high" for item in filename_findings
                )
                privacy["medium_findings"] = int(privacy["medium_findings"]) + sum(
                    item["severity"] == "medium" for item in filename_findings
                )
                privacy["classification"] = (
                    "restricted" if privacy["high_findings"] else "personal_data_review"
                )
            status = "ok" if loaded else "empty"
            message = ""
            if not loaded and path.suffix.lower() == ".pdf":
                status = "needs_ocr"
                message = "No searchable text was extracted; inspect the PDF and enable OCR if it is scanned."
            elif file_stats.get("total_pages", 0) > file_stats.get("pages_with_text", 0):
                status = "partial"
                message = "Some pages or slides yielded no searchable text; inspect them for scans or visual-only content."
            indexed = bool(loaded)
            if security["status"] == "quarantined":
                status = "quarantined"
                message = (
                    "High-risk instruction-like content was detected before normalization. "
                    "Review the security findings before allowing this file into the index."
                )
                indexed = not quarantine_suspicious
            elif privacy["classification"] == "restricted":
                status = "restricted"
                message = (
                    "Likely high-risk personal identifiers or credentials were detected. "
                    "The report contains categories and counts only; review the source file before indexing it."
                )
                indexed = not quarantine_sensitive
            elif security["status"] == "suspicious" and status == "ok":
                status = "suspicious"
                message = "Invisible or bidirectional Unicode was detected and normalized; review the findings."
            elif privacy["classification"] == "personal_data_review" and status == "ok":
                status = "privacy_review"
                message = (
                    "Possible contact information was detected. It remains indexed because it may be an official "
                    "public contact, but should be reviewed."
                )
            if indexed:
                units.extend(loaded)
            files.append({
                "source": source,
                "type": path.suffix.lower(),
                "status": status,
                "indexed": indexed,
                "sha256": file_hash,
                "bytes": file_size,
                "extracted_units": len(loaded),
                "pages_with_text": file_stats.get(
                    "pages_with_text", len({unit.page for unit in loaded if unit.page is not None})
                ),
                "total_pages": file_stats.get("total_pages"),
                "message": message,
                "security": security,
                "privacy": privacy,
            })
        except Exception as exc:
            failure = f"{source}: {exc}"
            failures.append(failure)
            files.append({
                "source": source,
                "type": path.suffix.lower(),
                "status": "error",
                "indexed": False,
                "sha256": file_hash,
                "bytes": file_size,
                "extracted_units": 0,
                "pages_with_text": 0,
                "total_pages": None,
                "message": str(exc),
                "security": {"status": "unknown", "findings": []},
                "privacy": {"classification": "unknown", "findings": []},
            })
    if not units:
        quarantined = [item["source"] for item in files if item["status"] in {"quarantined", "restricted"}]
        reason = "No safe searchable text remained after document screening."
        if quarantined:
            reason += " Quarantined: " + ", ".join(quarantined) + "."
        if failures:
            reason += " Extraction failures: " + "; ".join(failures)
        raise RuntimeError(reason)
    cleaned = []
    for unit in units:
        text = clean_text(unit.text)
        if text:
            cleaned.append(SourceUnit(text, unit.source, unit.page, unit.section, unit.document_type))
    cleaned = remove_repeated_page_lines(cleaned)
    if diagnostics is not None:
        diagnostics.update({
            "files": files,
            "ok": sum(item["status"] == "ok" for item in files),
            "indexed_files": sum(bool(item["indexed"]) for item in files),
            "quarantined": sum(item["status"] == "quarantined" for item in files),
            "restricted": sum(item["status"] == "restricted" for item in files),
            "needs_attention": sum(item["status"] != "ok" for item in files),
            "failures": failures,
        })
    if failures:
        print("Warnings while loading documents:")
        for failure in failures:
            print(f"- {failure}")
    return cleaned


def document_inventory(directory: Path) -> dict:
    paths = [path for path in directory.rglob("*") if path.is_file()]
    supported = [path for path in paths if path.suffix.lower() in SUPPORTED_EXTENSIONS]
    return {
        "total_files": len(paths),
        "supported_files": len(supported),
        "unsupported": sorted({path.suffix.lower() or "<no extension>" for path in paths if path not in supported}),
        "by_extension": {
            extension: sum(path.suffix.lower() == extension for path in supported)
            for extension in sorted({path.suffix.lower() for path in supported})
        },
    }
